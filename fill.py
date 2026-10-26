# fill.py - puts the daily numbers into YOUR template.
# It edits your original file directly, so your formulas, styles and layout stay the same.
# It also saves the RESULT of every formula, so totals show even when the file is only
# previewed (Telegram preview, phone, Excel "Protected View").
import copy
import re
import zipfile
from datetime import date, datetime
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

TEMPLATE = Path(__file__).parent / "template.xlsx"

# Which row each item is in your template (columns D, E, F = cash, card, aba)
ROOM_ROWS = {"pool": 2, "city": 3, "double_twin": 4, "pod": 5, "dorm4": 6, "dorm6": 7}
TOUR_ROWS = {"killing_field": 15, "bus": 16, "scooter": 17,
             "airport": 18, "sunset_boat": 19, "night_city": 20}
BAR_ROW = 24
EXPENSE_FIRST, EXPENSE_LAST = 30, 43  # A = name, B = amount

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
Q = lambda tag: f"{{{NS}}}{tag}"


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0


# ---------- 1. what to write ----------
def collect_inputs(data):
    cells, zero_cells = {}, set()

    def put_row(row, values):
        values = (list(values or []) + [0, 0, 0])[:3]
        if all(num(v) == 0 for v in values):   # nothing entered (or cleared) -> leave blank
            return
        for col, v in zip("DEF", values):      # entered: write all 3, so "20 0 0" shows 20, 0, 0
            cells[f"{col}{row}"] = num(v)
            if num(v) == 0:
                zero_cells.add(f"{col}{row}")

    for key, row in ROOM_ROWS.items():
        put_row(row, (data.get("rooms") or {}).get(key))
    for key, row in TOUR_ROWS.items():
        put_row(row, (data.get("tours") or {}).get(key))
    put_row(BAR_ROW, data.get("bar"))

    d = data.get("discount") or {}
    if num(d.get("amount")) != 0:
        cells["B10"] = num(d["amount"])
        if d.get("reason"):
            cells["B11"] = str(d["reason"])

    row = EXPENSE_FIRST
    for e in data.get("expenses") or []:
        if row > EXPENSE_LAST:
            break
        cells[f"A{row}"] = str(e.get("name", ""))
        cells[f"B{row}"] = num(e.get("amount"))
        row += 1
    return cells, zero_cells


# ---------- 2. calculate your formulas (SUM and cell links) ----------
def _cells_in(ref):
    ref = ref.replace("$", "")
    if ":" in ref:
        a, b = ref.split(":")
        ma, mb = re.match(r"([A-Z]+)(\d+)", a), re.match(r"([A-Z]+)(\d+)", b)
        for r in range(int(ma.group(2)), int(mb.group(2)) + 1):
            for c in range(column_index_from_string(ma.group(1)), column_index_from_string(mb.group(1)) + 1):
                yield f"{get_column_letter(c)}{r}"
    else:
        yield ref


def calculate(inputs):
    """Return {cell: value} for every formula cell in the template, or {} if a formula is unknown."""
    ws = load_workbook(TEMPLATE).worksheets[0]
    formulas, consts = {}, {}
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                formulas[c.coordinate] = c.value[1:].strip()
            elif isinstance(c.value, (int, float)):
                consts[c.coordinate] = c.value
    memo = {}

    def value(ref):
        if ref in inputs:
            v = inputs[ref]
            return v if isinstance(v, (int, float)) else 0
        if ref in formulas:
            if ref not in memo:
                memo[ref] = evaluate(formulas[ref])
            return memo[ref]
        return consts.get(ref, 0)

    def evaluate(expr):
        m = re.fullmatch(r"SUM\((.*)\)", expr, re.I)
        if m:
            return sum(value(r) for part in m.group(1).split(",") for r in _cells_in(part.strip()))
        if re.fullmatch(r"\$?[A-Z]+\$?\d+", expr):
            return value(expr.replace("$", ""))
        raise ValueError(expr)

    try:
        return {ref: value(ref) for ref in formulas}
    except ValueError:  # a formula we cannot calculate: Excel will do it when the file opens
        return {}


# ---------- 3. write into the original file ----------
def _get_cell(sheet_data, ref):
    m = re.match(r"([A-Z]+)(\d+)", ref)
    col, rown = column_index_from_string(m.group(1)), int(m.group(2))
    row = None
    for r in sheet_data.findall(Q("row")):
        rr = int(r.get("r"))
        if rr == rown:
            row = r
            break
        if rr > rown:
            row = etree.Element(Q("row"), r=str(rown))
            r.addprevious(row)
            break
    if row is None:
        row = etree.SubElement(sheet_data, Q("row"), r=str(rown))
    for c in row.findall(Q("c")):
        cc = column_index_from_string(re.match(r"[A-Z]+", c.get("r")).group(0))
        if cc == col:
            return c
        if cc > col:
            new = etree.Element(Q("c"), r=ref)
            c.addprevious(new)
            return new
    return etree.SubElement(row, Q("c"), r=ref)


def _clear(c):
    for child in list(c):
        if child.tag in (Q("v"), Q("is")):
            c.remove(child)
    c.attrib.pop("t", None)


def _fmt(v):
    return str(int(v)) if float(v).is_integer() else repr(float(v))


# Your input cells use the accounting format that shows a zero as "-".
# For zeros that were typed in, we use the same format but show 0.00 instead of "-".
ZERO_SHOWN_FORMAT = '_("$"* #,##0.00_);_("$"* \\(#,##0.00\\);_("$"* 0.00_);_(@_)'
DATE_FORMAT = "dd/mm/yyyy"
DATE_CELL = "B23"


class Styles:
    """Adds copies of existing cell styles that differ only in number format."""

    def __init__(self, xml_bytes):
        self.root = etree.fromstring(xml_bytes)
        self.numfmts = self.root.find(Q("numFmts"))
        self.xfs = self.root.find(Q("cellXfs"))
        self.cache = {}

    def _format_id(self, code):
        for nf in self.numfmts:
            if nf.get("formatCode") == code:
                return nf.get("numFmtId")
        new_id = str(max([163] + [int(n.get("numFmtId")) for n in self.numfmts]) + 1)
        etree.SubElement(self.numfmts, Q("numFmt"), numFmtId=new_id, formatCode=code)
        self.numfmts.set("count", str(len(self.numfmts)))
        return new_id

    def variant(self, base_s, code):
        key = (base_s, code)
        if key not in self.cache:
            xf = copy.deepcopy(self.xfs[base_s])
            xf.set("numFmtId", self._format_id(code))
            xf.set("applyNumberFormat", "1")
            self.xfs.append(xf)
            self.xfs.set("count", str(len(self.xfs)))
            self.cache[key] = len(self.xfs) - 1
        return self.cache[key]

    def tobytes(self):
        return etree.tostring(self.root, xml_declaration=True, encoding="UTF-8", standalone=True)


def fill_template(data, out_path, report_date=None):
    report_date = report_date or datetime.now().date()
    inputs, zero_cells = collect_inputs(data)
    results = calculate(inputs)

    with zipfile.ZipFile(TEMPLATE) as z:
        parts = {name: z.read(name) for name in z.namelist()}
        order = z.namelist()

    styles = Styles(parts["xl/styles.xml"])
    sheet_name = "xl/worksheets/sheet1.xml"
    root = etree.fromstring(parts[sheet_name])
    sheet_data = root.find(Q("sheetData"))

    for ref, v in inputs.items():                       # your typed numbers / text
        c = _get_cell(sheet_data, ref)
        _clear(c)
        for f in c.findall(Q("f")):
            c.remove(f)
        if isinstance(v, str):
            c.set("t", "inlineStr")
            t = etree.SubElement(etree.SubElement(c, Q("is")), Q("t"))
            t.text = v
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        else:
            if ref in zero_cells:                       # typed 0 -> show 0, not "-"
                c.set("s", str(styles.variant(int(c.get("s", 0)), ZERO_SHOWN_FORMAT)))
            etree.SubElement(c, Q("v")).text = _fmt(v)

    c = _get_cell(sheet_data, DATE_CELL)                # real date, shown day/month/year
    _clear(c)
    c.set("s", str(styles.variant(int(c.get("s", 0)), DATE_FORMAT)))
    serial = (report_date - date(1899, 12, 30)).days
    etree.SubElement(c, Q("v")).text = str(serial)

    for ref, v in results.items():                      # saved result of each formula
        c = _get_cell(sheet_data, ref)
        _clear(c)
        etree.SubElement(c, Q("v")).text = _fmt(v)

    for r in sheet_data.findall(Q("row")):              # row "spans" hints may be outdated
        r.attrib.pop("spans", None)
    parts[sheet_name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    parts["xl/styles.xml"] = styles.tobytes()

    wbx = parts["xl/workbook.xml"].decode("utf-8")      # Excel also recalculates on open
    if "<calcPr" in wbx and "fullCalcOnLoad" not in wbx:
        wbx = re.sub(r"<calcPr", '<calcPr fullCalcOnLoad="1"', wbx, count=1)
    parts["xl/workbook.xml"] = wbx.encode("utf-8")

    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as z:
        for name in order:
            z.writestr(name, parts[name])


def summarize(data):
    """Totals for the Telegram message (same logic as your formulas)."""
    def s(v):
        return sum(num(x) for x in (v or []))
    rooms = sum(s(v) for v in (data.get("rooms") or {}).values())
    tours = sum(s(v) for v in (data.get("tours") or {}).values())
    bar = s(data.get("bar"))
    expenses = sum(num(e.get("amount")) for e in (data.get("expenses") or []))
    return {"rooms": rooms, "tours": tours, "bar": bar,
            "total": rooms + tours + bar, "expenses": expenses}